"""Export layer — CSV, JSON, and Excel (.xlsx) exports.

Concept note — why multiple formats: CSV for simplicity, JSON for
interoperability, Excel for human review with multiple sheets.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .db import Database
from .logging import get_logger

logger = get_logger(__name__)


def export_csv(db: Database, path: str | Path) -> None:
    """Export results to CSV."""
    conn = db.connect()
    rows = conn.execute(
        """
        SELECT c.name, c.official_domain, d.status, d.destination_type,
               d.careers_url, d.job_board_url, d.ats_provider, d.ats_slug,
               d.open_jobs_count, d.discovery_method, d.confidence
        FROM companies c
        JOIN career_destinations d ON c.id = d.company_id
        """
    ).fetchall()

    with open(path, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([
            "name", "domain", "status", "destination_type",
            "careers_url", "job_board_url", "ats_provider", "ats_slug",
            "open_jobs", "discovery_method", "confidence",
        ])
        for row in rows:
            writer.writerow(row)

    logger.info(f"exported {len(rows)} rows to {path}")


def export_json(db: Database, path: str | Path) -> None:
    """Export results to JSON."""
    conn = db.connect()
    rows = conn.execute(
        """
        SELECT c.name, c.official_domain, d.status, d.destination_type,
               d.careers_url, d.job_board_url, d.ats_provider, d.ats_slug,
               d.open_jobs_count, d.discovery_method, d.confidence
        FROM companies c
        JOIN career_destinations d ON c.id = d.company_id
        """
    ).fetchall()

    data = [
        {
            "name": r[0],
            "domain": r[1],
            "status": r[2],
            "destination_type": r[3],
            "careers_url": r[4],
            "job_board_url": r[5],
            "ats_provider": r[6],
            "ats_slug": r[7],
            "open_jobs": r[8],
            "discovery_method": r[9],
            "confidence": r[10],
        }
        for r in rows
    ]

    with open(path, "w") as f:
        json.dump(data, f, indent=2)

    logger.info(f"exported {len(data)} records to {path}")


def export_xlsx(db: Database, path: str | Path) -> None:
    """Export results to Excel with a summary sheet and a sheet per status."""
    conn = db.connect()

    wb = Workbook()

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Status", "Count"])
    rows = conn.execute(
        "SELECT status, COUNT(*) FROM career_destinations GROUP BY status"
    ).fetchall()
    for r in rows:
        ws.append([r[0], r[1]])

    # Style header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")

    # Sheet per status
    for status in ["verified", "needs_review", "not_found", "blocked", "inactive"]:
        rows = conn.execute(
            """
            SELECT c.name, c.official_domain, d.destination_type,
                   d.careers_url, d.job_board_url, d.ats_provider, d.ats_slug,
                   d.open_jobs_count, d.discovery_method, d.confidence
            FROM companies c
            JOIN career_destinations d ON c.id = d.company_id
            WHERE d.status = ?
            """,
            (status,),
        ).fetchall()

        if not rows:
            continue

        ws = wb.create_sheet(title=status)
        ws.append([
            "name", "domain", "destination_type", "careers_url",
            "job_board_url", "ats_provider", "ats_slug",
            "open_jobs", "discovery_method", "confidence",
        ])
        for r in rows:
            ws.append(list(r))

        # Style header
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")

    wb.save(path)
    logger.info(f"exported to {path}")
